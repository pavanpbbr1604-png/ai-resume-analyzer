import docx
import pytest
from pathlib import Path
from app.document.docx_parser import parse_docx, convert_docx_to_html

@pytest.fixture
def sample_docx(tmp_path) -> Path:
    doc_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("WORK EXPERIENCE", level=1)
    
    p = doc.add_paragraph("Developed scalable web microservices using Python and Flask.")
    run = p.runs[0]
    run.font.name = "Calibri"
    run.font.size = docx.shared.Pt(11)
    run.bold = True
    
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph("Python, JavaScript, SQL")
    
    doc.save(doc_path)
    return doc_path

def test_parse_docx_sections_and_runs(sample_docx):
    norm_doc = parse_docx(sample_docx, document_id="doc_test_123")
    assert norm_doc.document_id == "doc_test_123"
    assert len(norm_doc.sections) >= 2
    
    # Check section detection
    exp_sec = next((s for s in norm_doc.sections if s.section_type == "EXPERIENCE"), None)
    assert exp_sec is not None
    assert len(exp_sec.paragraphs) > 0
    
    body_p = next(p for p in exp_sec.paragraphs if "Developed" in p.full_text)
    assert "Developed scalable" in body_p.full_text
    assert body_p.runs[0].formatting.bold is True
    assert body_p.runs[0].formatting.font_family == "Calibri"

def test_convert_docx_to_html(sample_docx):
    html = convert_docx_to_html(sample_docx)
    assert "<h1" in html or "<p" in html
    assert "WORK EXPERIENCE" in html
