import docx
import pytest
from pathlib import Path
from app.document.inheritance_engine import FormattingInheritanceEngine, ConcurrencyConflictException

@pytest.fixture
def docx_file(tmp_path) -> Path:
    p_file = tmp_path / "test_edit.docx"
    doc = docx.Document()
    p = doc.add_paragraph()
    r = p.add_run("Developed a website using Python")
    r.font.name = "Arial"
    r.font.size = docx.shared.Pt(12)
    r.bold = True
    doc.save(p_file)
    return p_file

def test_apply_replacement_formatting_preserved(docx_file):
    engine = FormattingInheritanceEngine()
    success = engine.update_docx_file(
        file_path=docx_file,
        paragraph_index=1,
        original_text="Developed a website using Python",
        suggested_text="Engineered a website using Python.",
    )
    assert success is True
    
    # Verify resulting document text and formatting
    res_doc = docx.Document(docx_file)
    p = res_doc.paragraphs[0]
    assert p.text == "Engineered a website using Python."
    assert p.runs[0].font.name == "Arial"
    assert p.runs[0].bold is True

def test_concurrency_conflict_exception(docx_file):
    engine = FormattingInheritanceEngine()
    doc = docx.Document(docx_file)
    p = doc.paragraphs[0]
    
    with pytest.raises(ConcurrencyConflictException):
        engine.apply_replacement(
            paragraph=p,
            original_text="Developed a website using Python",
            suggested_text="New text",
            expected_hash="invalid_hash_xyz",
        )
